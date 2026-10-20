"""The loaders `instrument-index` and `equipment-list` (P1-I8-T07b). Provided; do not edit."""

from __future__ import annotations

import csv
from importlib.metadata import entry_points
from io import BytesIO

import pytest
from openpyxl import Workbook
from tl_mod_engineering.imports import mapping as tags_mapping
from tl_mod_legacy.errors import SourceFileError
from tl_mod_legacy.loaders.equipment_list import loader as equipment_loader
from tl_mod_legacy.loaders.instrument_index import loader as instrument_loader
from tl_mod_legacy.model import use_loaders
from tl_mod_legacy.targets import comma_mapping_text
from tl_mod_legacy.testing import ADMIN, PROJECT, TAG_FILE, LegacyWorld

EQUIPMENT_HEADER = [
    "EQUIP NO",
    "DESCRIPTION",
    "TYPE",
    "DESIGN PRESSURE (BAR)",
    "DESIGN TEMP (C)",
    "WEIGHT (KG)",
    "VENDOR",
    "REMARKS",
]
EQUIPMENT_ROWS = [
    ["47-P-0101", "Feed pump A", "PUMP", "42", "120", "1800", "Borealis", ""],
    ["47-P-0102", "Feed pump B", "PUMP", "42", "120", "1800", "Borealis", "spare"],
    ["47-V-0201", "Feed surge drum", "VESSEL", "10", "80", "9500", "Acme Vessels", ""],
    ["47-E-0301", "Feed effluent exchanger", "EXCHANGER", "35", "300", "7200", "Acme HX", ""],
    ["47-K-0401", "Recycle compressor", "COMPRESSOR", "60", "150", "21000", "Borealis", ""],
    ["47-M-0501", "Feed pump A motor", "MOTOR", "", "", "950", "Borealis", ""],
    ["", "A row with no number", "PUMP", "", "", "", "", ""],
    ["47-V-0202", "A boiler", "BOILER", "", "", "", "", ""],
    ["47-P-0101", "Feed pump A again", "PUMP", "", "", "", "", ""],
]


def workbook(rows: list[list[str]], *, header: list[str], sheet: str, titles: int = 0) -> bytes:
    """A workbook with a cover sheet, then ``sheet`` with ``titles`` title rows above the header."""
    book = Workbook()
    cover = book.active
    assert cover is not None
    cover.title = "Cover"
    cover.append(["Exported from the incumbent system"])
    data = book.create_sheet(sheet)
    for n in range(titles):
        data.append([f"Title row {n + 1}"])
    data.append(header)
    for row in rows:
        data.append(row)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def seed_rows() -> tuple[list[str], list[list[str]]]:
    lines = list(csv.reader(TAG_FILE.read_text(encoding="utf-8").splitlines()))
    return lines[0], lines[1:]


def test_the_loaders_describe_what_they_read_and_write() -> None:
    first, second = instrument_loader(), equipment_loader()
    assert (first.name, first.system, first.table, first.target) == (
        "instrument-index",
        "instr",
        "index",
        "tags",
    )
    assert first.file_types == ("xlsx", "csv") and first.sheet is None and first.header_row == 1
    assert (second.name, second.system, second.table, second.target) == (
        "equipment-list",
        "equip",
        "list",
        "tags",
    )
    assert second.file_types == ("xlsx",)
    assert second.sheet == "Equipment List" and second.header_row == 3
    assert first.confidentiality is None and second.confidentiality is None
    assert first.description == "Instrument index of the incumbent system (workbook or CSV)."
    assert second.description == "Equipment list of the incumbent system (workbook)."


def test_the_templates_name_their_columns() -> None:
    first = tags_mapping.parse_mapping(comma_mapping_text(instrument_loader()), source="a")
    second = tags_mapping.parse_mapping(comma_mapping_text(equipment_loader()), source="b")
    assert (first.name, second.name) == ("instrument-index", "equipment-list")
    assert [c.column for c in second.columns] == ["EQUIP NO", "DESCRIPTION", "TYPE"]
    assert second.columns[2].values["EXCHANGER"] == "HeatExchanger"
    assert second.ignore == [
        "DESIGN PRESSURE (BAR)",
        "DESIGN TEMP (C)",
        "WEIGHT (KG)",
        "VENDOR",
        "REMARKS",
    ]
    assert "TAG NO" in [c.column for c in first.columns]


def test_they_are_installed_under_their_names() -> None:
    found = {e.name: e for e in entry_points(group="throughline.loaders")}
    assert found["instrument-index"].value == "tl_mod_legacy.loaders.instrument_index:loader"
    assert found["equipment-list"].value == "tl_mod_legacy.loaders.equipment_list:loader"
    assert found["instrument-index"].load()() == instrument_loader()
    assert found["equipment-list"].load()() == equipment_loader()


def test_the_instrument_index_loads_from_a_workbook_like_from_a_file(world: LegacyWorld) -> None:
    header, rows = seed_rows()
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    out = BytesIO()
    book.save(out)
    with use_loaders({"instrument-index": instrument_loader()}):
        run = world.through_plan("instrument-index", "index.xlsx", out.getvalue())
        assert (run.key, run.plan_new, run.plan_existing, run.plan_invalid) == (
            "instr/index#1",
            11,
            0,
            3,
        )
        outcome = world.service.load(PROJECT, run.key, actor=ADMIN)
        assert (outcome.created, outcome.skipped_invalid, outcome.source_layer_rows) == (11, 3, 11)
        done, report = world.service.reconcile(PROJECT, run.key, actor=ADMIN)
        assert done.mismatches == 0 and report["matched"] == 11
        # the same rows as a CSV file: every tag is there already
        from_file = world.through_plan("instrument-index", "index.csv", TAG_FILE.read_bytes())
    assert (from_file.plan_new, from_file.plan_existing, from_file.plan_invalid) == (0, 11, 3)


def test_the_equipment_list_is_read_from_its_sheet_below_two_title_rows(
    world: LegacyWorld,
) -> None:
    book = workbook(EQUIPMENT_ROWS, header=EQUIPMENT_HEADER, sheet="Equipment List", titles=2)
    with use_loaders({"equipment-list": equipment_loader()}):
        run = world.through_plan("equipment-list", "equipment.xlsx", book)
        assert run.row_count == 9 and run.column_count == 8
        assert (run.key, run.plan_new, run.plan_existing, run.plan_invalid) == (
            "equip/list#1",
            6,
            0,
            3,
        )
        outcome = world.service.load(PROJECT, run.key, actor=ADMIN)
        assert (outcome.created, outcome.skipped_invalid) == (6, 3)
        _, report = world.service.reconcile(PROJECT, run.key, actor=ADMIN)
    assert report["mismatches"] == 0 and report["matched"] == 6 and report["invalid_rows"] == 3
    classes = {
        r["key"]: r["class_code"]
        for r in world.rows(
            "SELECT r.key, t.class_code FROM cur_engineering_tag t "
            "JOIN cur_core_record r ON r.id = t.id"
        )
    }
    assert classes == {
        "47-P-0101": "Pump",
        "47-P-0102": "Pump",
        "47-V-0201": "Vessel",
        "47-E-0301": "HeatExchanger",
        "47-K-0401": "Compressor",
        "47-M-0501": "Motor",
    }


def test_the_design_data_is_kept_in_the_source_layer(world: LegacyWorld) -> None:
    book = workbook(EQUIPMENT_ROWS, header=EQUIPMENT_HEADER, sheet="Equipment List", titles=2)
    with use_loaders({"equipment-list": equipment_loader()}):
        world.loaded("equipment-list", "equipment.xlsx", book)
    tag = world.rows("SELECT id FROM cur_core_record WHERE key = :k", k="47-P-0101")[0]["id"]
    values = {
        v["path"]: v["value_text"]
        for v in world.rows(
            "SELECT path, value_text FROM cur_pset_values "
            "WHERE record_id = :i AND layer = 'source'",
            i=tag,
        )
    }
    assert any(path.endswith("design_pressure_bar") for path in values)
    assert any(path.endswith("vendor") for path in values)
    assert not any(path.endswith("remarks") for path in values)  # an empty cell is not kept


def test_a_workbook_without_the_sheet_is_refused_and_the_run_stays_open(
    world: LegacyWorld,
) -> None:
    book = workbook(EQUIPMENT_ROWS, header=EQUIPMENT_HEADER, sheet="Other Sheet", titles=2)
    with use_loaders({"equipment-list": equipment_loader()}):
        run = world.open("equipment-list", "equipment.xlsx", book)
        with pytest.raises(SourceFileError):
            world.service.extract(PROJECT, run.key, actor=ADMIN)
    assert world.run(run.key).state == "Opened"
    assert "LoadRun.Failed" in world.events_of(run.id)


def test_the_equipment_list_does_not_take_a_csv_file(world: LegacyWorld) -> None:
    with use_loaders({"equipment-list": equipment_loader()}):
        with pytest.raises(SourceFileError, match="reads xlsx"):
            world.open("equipment-list", "equipment.csv", b"EQUIP NO\n1\n")
