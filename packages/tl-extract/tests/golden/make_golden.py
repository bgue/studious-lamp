"""The golden dataset of the extractors (P1-I3B-T12): three sources, built the same way every time.

    uv run python packages/tl-extract/tests/golden/make_golden.py --check
    uv run python packages/tl-extract/tests/golden/make_golden.py --write-sources
    uv run python packages/tl-extract/tests/golden/make_golden.py --write-expected

The sources are a PDF, a CSV and an XLSX, written by this script from fixed text and fixed times
(no clock, no random numbers, no download). ``--check`` rebuilds them, compares them with the
committed files, runs the matching extractor over each and compares the result with the committed
expected JSON. ``--write-expected`` rewrites the expected JSON from the extractors: read the diff
before committing it, the point of the files is that a person has judged them.

The XLSX is stored (not compressed) with fixed entry times, and the modified date that openpyxl
stamps on save is replaced, so that its bytes depend neither on the zlib that happens to be
installed nor on the time of the run.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))  # extractsupport lives in tests/

from extractsupport import make_pdf  # noqa: E402
from tl_core.derivation import Extractor, SourceFile  # noqa: E402

SOURCES = HERE / "sources"
EXPECTED = HERE / "expected"
FIXED_TIME = dt.datetime(2024, 3, 1, 10, 15, 30)

HEADER_MAP = {
    "Tag No.": "tag_number",
    "Service": "service",
    "Class": "tag_class",
    "Size (in)": "size_in",
    "Material": "body_material",
    "MFR": "manufacturer",
    "Fail": "fail_action",
    "Remarks": "remarks",
}


@dataclass(frozen=True)
class Case:
    name: str
    filename: str
    content_type: str
    extractor: str
    options: dict[str, Any]


CASES = [
    Case("pdf", "ISD-001.pdf", "application/pdf", "pdf", {}),
    Case(
        "csv",
        "instrument-index-export.csv",
        "text/csv",
        "tabular",
        {"header_map": HEADER_MAP},
    ),
    Case(
        "xlsx",
        "instrument-index-book.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "tabular",
        {"sheet": "Index", "header_row": 3, "header_map": HEADER_MAP},
    ),
]


def build_pdf() -> bytes:
    """Three blank pages and an information dictionary with a UTF-16 title, both dates and a
    producer. The title is a hex string, which is how PDF writers store non-ASCII text."""
    title = "﻿Instrument index – ISD-001".encode("utf-16-be").hex().upper()
    return make_pdf(
        {
            "Title": f"<{title}>".encode(),
            "Author": "Acme Instrumentation",
            "Producer": "Throughline golden script",
            "CreationDate": "D:20240301101530+02'00'",
            "ModDate": "D:20240305170000Z",
        },
        pages=3,
    )


CSV_TEXT = (
    "Tag No.;Service;Class;Size (in);Material;MFR;Fail;Remarks\n"
    "47-FV-1001;Feed gas flow control;ControlValve;4;SS316L;Acme Valves;FC;\n"
    "47-FV-1002;Fuel gas flow control;ControlValve;6;CS;Acme Valves;FO;Replaces FV-0902\n"
    '47-PV-1003;"Flare header; north leg, pressure control";ControlValve;8;SS316L;Borealis;FO;\n'
    "47-HV-1004;  Feed gas isolation  ;ManualValve;4;CS;Borealis;;\n"
    "\n"
    "47-PSV-1005;Separator relief;SafetyValve;3;SS316L;Borealis;;Set 18 barg\n"
    "47-FT-1006;Feed gas flow transmitter;Transmitter;;;;;\n"
    '47-LIC-1007;"Separator level\ncontroller";Controller;;;;;"Shown on PID-4703"\n'
    "47-P-1008;Condensate pump;Pump;;;;;Bell\x07 rang\n"
    "47-V-1009;Feed separator;Vessel;;;;;\n"
    ";;;;;;;\n"
    "47-E-1010;Reactor effluent cooler;HeatExchanger;;;;;;stray cell\n"
)


def build_csv() -> bytes:
    """A semicolon export with a byte-order mark, quoted cells (a delimiter inside, a line break
    inside), a blank line, a line of empty cells, a control character and a stray cell."""
    return b"\xef\xbb\xbf" + CSV_TEXT.encode("utf-8")


def build_xlsx() -> bytes:
    """A workbook of three sheets: a cover, the index below a two-line title block, and a hidden
    archive. The index holds numbers, a date, a bool, a blank row and a formula (no stored value,
    so it reads as an empty cell)."""
    book = openpyxl.Workbook()
    cover = book.active
    assert cover is not None
    cover.title = "Cover"
    cover.append(["Instrument index ISD-001 revision A"])
    index = book.create_sheet("Index")
    index.append(["Acme Hydrotreater, Unit 47"])
    index.append(["Issued for review"])
    index.append(list(HEADER_MAP) + ["Issued", "Witnessed"])
    index.append(
        ["47-FV-1001", "Feed gas flow control", "ControlValve", 4, "SS316L", "Acme Valves", "FC"]
        + [None, dt.date(2024, 3, 1), True]
    )
    index.append(
        ["47-FV-1002", "Fuel gas flow control", "ControlValve", 6.5, "CS", "Acme Valves", "FO"]
        + ["Replaces FV-0902", dt.datetime(2024, 3, 1, 13, 5, 0), False]
    )
    index.append([])
    index.append(["47-HV-1004", "Feed gas isolation", "ManualValve", "=D4+D5", "CS", "Borealis"])
    index.append(["47-FT-1006", "Feed gas flow transmitter", "Transmitter"])
    archive = book.create_sheet("Archive")
    archive.append(["old"])
    archive.sheet_state = "hidden"
    book.properties.creator = "Acme Instrumentation"
    book.properties.title = "Instrument index"
    book.properties.created = FIXED_TIME
    book.properties.modified = FIXED_TIME
    plain = io.BytesIO()
    book.save(plain)
    out = io.BytesIO()
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2024, 3, 1, 10, 15, 30))
            info.compress_type = zipfile.ZIP_STORED
            info.external_attr = 0o600 << 16
            data = source.read(name)
            if name == "docProps/core.xml":
                # openpyxl stamps the time of the save into the modified date
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2024-03-05T17:00:00Z\g<2>",
                    data,
                )
            target.writestr(info, data)
    return out.getvalue()


BUILDERS = {"pdf": build_pdf, "csv": build_csv, "xlsx": build_xlsx}


def extractors() -> dict[str, Extractor]:
    from tl_extract.ocr import OcrExtractor
    from tl_extract.pdf import PdfExtractor
    from tl_extract.tabular import TabularExtractor

    return {"tabular": TabularExtractor(), "pdf": PdfExtractor(), "ocr": OcrExtractor()}


def run_case(case: Case, data: bytes) -> dict[str, Any]:
    """The extraction result of ``data`` for ``case`` as plain JSON-able data."""
    source = SourceFile(
        file_id="f-golden",
        filename=case.filename,
        content_type=case.content_type,
        size=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        source_ref="src.doc.ISD-001@A",
    )
    extractor = extractors()[case.extractor]
    assert extractor.accepts(source), case.name
    result = extractor.extract(source, io.BytesIO(data), dict(case.options))
    return result.model_dump(mode="json")


def source_path(case: Case) -> Path:
    return SOURCES / case.filename


def expected_path(case: Case) -> Path:
    return EXPECTED / f"{case.name}.json"


def dumps(data: dict[str, Any]) -> str:
    return json.dumps(data, indent=2, ensure_ascii=False, sort_keys=True) + "\n"


def main(argv: list[str]) -> int:
    mode = argv[0] if argv else "--check"
    if mode == "--write-sources":
        SOURCES.mkdir(exist_ok=True)
        for case in CASES:
            source_path(case).write_bytes(BUILDERS[case.name]())
        return 0
    if mode == "--write-expected":
        EXPECTED.mkdir(exist_ok=True)
        for case in CASES:
            expected_path(case).write_text(
                dumps(run_case(case, source_path(case).read_bytes())), encoding="utf-8"
            )
        return 0
    if mode != "--check":
        print("usage: make_golden.py [--check | --write-sources | --write-expected]")
        return 2
    problems: list[str] = []
    for case in CASES:
        committed = source_path(case).read_bytes()
        if BUILDERS[case.name]() != committed:
            problems.append(f"{case.filename}: the script builds different bytes")
        if dumps(run_case(case, committed)) != expected_path(case).read_text(encoding="utf-8"):
            problems.append(f"{case.name}: the extractor's result differs from the expected file")
    for line in problems:
        print(line)
    print("golden dataset ok" if not problems else f"{len(problems)} problem(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
